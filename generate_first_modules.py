#!/usr/bin/env python3
"""
Generate test cases for first 4 modules (New Chat, Saved Prompts, History, Connect)
Format: TC-001, TC-002... per module (resets for each new sheet)
Columns: ID | Title | Description | Precondition | Severity | Steps | Expected | Actual
"""

import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Module configuration
MODULES = [
    {
        "name": "New Chat",
        "file": "tests/01-new-chat.spec.js",
        "tests": [
            ("Chat Interface Loads", "Verify new chat interface loads", "User logged in",
             "1. Navigate to new chat page\n2. Wait for interface to load\n3. Verify chat box visible\n4. Check input field ready",
             "Chat interface displays, ready for input", "PASS - Verified"),

            ("Send Message", "User can send message to AI", "At chat page with input ready",
             "1. Click chat input field\n2. Type test message\n3. Press send/enter\n4. Verify message appears in chat",
             "Message sent and appears in conversation", "Pending - Manual"),

            ("AI Response", "AI responds to user message", "Message sent to AI",
             "1. Wait for AI response\n2. Verify response appears below user message\n3. Check response is readable\n4. Verify formatting intact",
             "AI response displays correctly", "Pending - Manual"),

            ("Chat History Navigation", "Previous chats visible in sidebar", "At chat page",
             "1. Look for chat history/sidebar\n2. Verify previous chats listed\n3. Click on previous chat\n4. Verify chat loads",
             "Previous chats load when selected", "Pending - Manual"),

            ("Clear Chat", "User can clear conversation", "Active chat exists",
             "1. Look for clear/new chat option\n2. Click to clear\n3. Confirm if prompted\n4. Verify chat cleared",
             "Chat history cleared, new conversation starts", "Pending - Manual"),
        ]
    },
    {
        "name": "Saved Prompts",
        "file": "tests/02-saved-prompts.spec.js",
        "tests": [
            ("Prompts Page Loads", "Saved prompts page displays", "User logged in",
             "1. Navigate to saved prompts section\n2. Wait for page load\n3. Verify prompt list visible\n4. Check interface ready",
             "Prompts page displays with list of saved items", "PASS - Verified"),

            ("Create Prompt", "User can create and save new prompt", "At saved prompts page",
             "1. Click 'add new prompt' button\n2. Enter prompt title\n3. Enter prompt content\n4. Click save\n5. Verify prompt appears in list",
             "New prompt created and visible in list", "Pending - Manual"),

            ("Edit Prompt", "User can modify saved prompt", "Prompt exists in list",
             "1. Click on existing prompt\n2. Click edit button\n3. Modify content\n4. Save changes\n5. Verify updates appear",
             "Prompt modifications saved and displayed", "Pending - Manual"),

            ("Delete Prompt", "User can remove prompt from list", "Prompt exists",
             "1. Click on prompt\n2. Click delete button\n3. Confirm deletion if prompted\n4. Verify prompt removed from list",
             "Prompt deleted successfully", "Pending - Manual"),

            ("Use Prompt", "User can use saved prompt in chat", "Saved prompt exists",
             "1. Open chat interface\n2. Click on saved prompt option\n3. Select prompt from list\n4. Verify prompt loads in chat\n5. Send to AI",
             "Prompt loads and can be sent to AI", "Pending - Manual"),
        ]
    },
    {
        "name": "History",
        "file": "tests/03-history.spec.js",
        "tests": [
            ("History Page Loads", "Chat history page displays", "User logged in",
             "1. Navigate to history/conversations section\n2. Wait for page load\n3. Verify conversation list visible\n4. Check dates and previews shown",
             "History page displays with all conversations", "PASS - Verified"),

            ("View Conversation", "User can open previous conversation", "History page loaded",
             "1. Click on conversation in history\n2. Wait for conversation to load\n3. Verify all messages visible\n4. Check timestamps present",
             "Previous conversation loads completely", "Pending - Manual"),

            ("Search History", "User can search conversation history", "At history page",
             "1. Click search field\n2. Type search keyword\n3. Wait for results to filter\n4. Verify matching conversations shown",
             "History filtered by search term", "Pending - Manual"),

            ("Clear History", "User can delete all history", "History exists",
             "1. Look for clear/delete all option\n2. Click clear all\n3. Confirm deletion in dialog\n4. Wait for clear to complete\n5. Verify history empty",
             "All history cleared successfully", "Pending - Manual"),

            ("Export History", "User can export conversation history", "Conversations exist",
             "1. Click export option\n2. Select format (PDF/text/etc)\n3. Verify download starts\n4. Check file format correct",
             "History exported in selected format", "Pending - Manual"),
        ]
    },
    {
        "name": "Connect",
        "file": "tests/04-connect.spec.js",
        "tests": [
            ("Integrations Page Loads", "Connect/integrations page displays", "User logged in",
             "1. Navigate to connect/integrations section\n2. Wait for page load\n3. Verify available integrations visible\n4. Check connection status shown",
             "Integrations page displays available services", "PASS - Verified"),

            ("View Integration", "User can view integration details", "At integrations page",
             "1. Click on integration card\n2. Wait for details to load\n3. Verify description and options visible\n4. Check status indicator present",
             "Integration details display correctly", "Pending - Manual"),

            ("Connect Integration", "User can authorize new integration", "At integration details",
             "1. Click 'connect' or 'authorize' button\n2. Follow authorization flow\n3. Complete any required login\n4. Verify connection successful\n5. Check status updates",
             "Integration connected and status shows active", "Pending - Manual"),

            ("Disconnect Integration", "User can disconnect integration", "Connected integration exists",
             "1. Click on connected integration\n2. Look for disconnect/remove option\n3. Click disconnect\n4. Confirm if prompted\n5. Verify status updates",
             "Integration disconnected, status shows inactive", "Pending - Manual"),

            ("Test Integration", "User can test integration connection", "Integration connected",
             "1. Click test/verify button\n2. Wait for test to run\n3. Verify test result displays\n4. Check connection health shown",
             "Integration test results display", "Pending - Manual"),
        ]
    }
]

def create_excel():
    """Create Excel with first 4 modules"""
    wb = Workbook()
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)

    severity_fill = {
        "HIGH": PatternFill(start_color="C00000", end_color="C00000", fill_type="solid"),
        "MEDIUM": PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid"),
        "LOW": PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid"),
    }
    severity_font = {k: Font(bold=True, color="FFFFFF") for k in severity_fill}

    border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    total_tests = 0

    # Create sheet for each module
    for module in MODULES:
        ws = wb.create_sheet(title=module["name"][:30])

        # Headers: ID | Title | Description | Precondition | Severity | Steps | Expected | Actual
        headers = ["Test ID", "Title", "Description", "Precondition", "Severity", "Steps to Reproduce", "Expected Result", "Actual Result"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Add test rows with TC-001 numbering per module
        for row_idx, test in enumerate(module["tests"], 2):
            test_id = f"TC-{row_idx-1:03d}"  # TC-001, TC-002, etc.
            title, desc, pre, steps, exp, actual = test

            # Determine severity
            if "delete" in title.lower() or "disconnect" in title.lower():
                severity = "HIGH"
            elif "edit" in title.lower() or "modify" in title.lower():
                severity = "MEDIUM"
            else:
                severity = "LOW"

            # Write cells
            ws.cell(row=row_idx, column=1, value=test_id)
            ws.cell(row=row_idx, column=2, value=title)
            ws.cell(row=row_idx, column=3, value=desc)
            ws.cell(row=row_idx, column=4, value=pre)
            ws.cell(row=row_idx, column=5, value=severity)
            ws.cell(row=row_idx, column=6, value=steps)
            ws.cell(row=row_idx, column=7, value=exp)
            ws.cell(row=row_idx, column=8, value=actual)

            # Format cells
            for col in range(1, 9):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="left")

                # Color severity column
                if col == 5:
                    cell.fill = severity_fill.get(severity)
                    cell.font = severity_font.get(severity)

        # Column widths
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 22
        ws.column_dimensions['C'].width = 35
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 50
        ws.column_dimensions['G'].width = 35
        ws.column_dimensions['H'].width = 35

        # Row heights
        ws.row_dimensions[1].height = 35
        for row in range(2, len(module["tests"]) + 2):
            ws.row_dimensions[row].height = 60

        total_tests += len(module["tests"])
        print(f"[OK] {module['name']}: {len(module['tests'])} test cases added (TC-001 to TC-{len(module['tests']):03d})")

    # Save workbook
    output_file = "C:\\Users\\dell\\Desktop\\M32.ai-AUtomation\\Project_Test_Cases_First_4_Modules.xlsx"
    wb.save(output_file)

    print(f"\n[SUCCESS] Excel file created!")
    print(f"[FILE] {output_file}")
    print(f"[TOTAL] {total_tests} professional test cases")
    print(f"\nFormat:")
    print(f"  Columns: ID | Title | Description | Precondition | Severity | Steps | Expected | Actual")
    print(f"  Each module resets: TC-001, TC-002, TC-003...")
    print(f"  Sheets: New Chat, Saved Prompts, History, Connect")

if __name__ == "__main__":
    create_excel()
